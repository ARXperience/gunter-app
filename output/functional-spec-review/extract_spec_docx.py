from __future__ import annotations

import json
import sys
from pathlib import Path

from docx import Document
from docx.document import Document as DocumentObject
from docx.table import Table
from docx.text.paragraph import Paragraph
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P


def iter_blocks(parent: DocumentObject):
    for child in parent.element.body.iterchildren():
        if isinstance(child, CT_P):
            yield Paragraph(child, parent)
        elif isinstance(child, CT_Tbl):
            yield Table(child, parent)


def normalize(value: str) -> str:
    return " ".join((value or "").replace("\xa0", " ").split())


def main() -> None:
    input_path = Path(sys.argv[1])
    output_path = Path(sys.argv[2])
    document = Document(input_path)
    lines: list[str] = []
    block_index = 0
    table_index = 0

    lines.append(f"SOURCE: {input_path}")
    lines.append(f"PARAGRAPHS: {len(document.paragraphs)}")
    lines.append(f"TABLES: {len(document.tables)}")
    lines.append("")

    for block in iter_blocks(document):
        block_index += 1
        if isinstance(block, Paragraph):
            text = normalize(block.text)
            if not text:
                continue
            style = normalize(block.style.name if block.style else "") or "Normal"
            lines.append(f"P{block_index:04d} [{style}] {text}")
            continue

        table_index += 1
        lines.append(f"T{table_index:02d} BEGIN")
        for row_index, row in enumerate(block.rows, start=1):
            cells = [normalize(cell.text) for cell in row.cells]
            lines.append(f"T{table_index:02d}R{row_index:02d} | " + " | ".join(cells))
        lines.append(f"T{table_index:02d} END")
        lines.append("")

    output_path.write_text("\n".join(lines), encoding="utf-8")
    summary = {
        "paragraphs": len(document.paragraphs),
        "tables": len(document.tables),
        "blocks": block_index,
        "characters": sum(len(line) for line in lines),
        "output": str(output_path),
    }
    print(json.dumps(summary, ensure_ascii=False))


if __name__ == "__main__":
    main()
